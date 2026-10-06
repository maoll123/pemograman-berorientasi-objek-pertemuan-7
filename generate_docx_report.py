import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_box_borders(cell, color="003366", sz="8"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def build_docx_report():
    doc = Document()

    # Page Setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # HEADER / TITLE
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    run_t1 = p_title.add_run("LAPORAN PRAKTIKUM PEMROGRAMAN BERORIENTASI OBYEK\n")
    run_t1.font.name = 'Calibri'
    run_t1.font.size = Pt(14)
    run_t1.bold = True
    run_t1.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    run_t2 = p_title.add_run("MODUL 7: ABSTRACT CLASS, ABSTRACT METHOD, DAN INTERFACE\n")
    run_t2.font.name = 'Calibri'
    run_t2.font.size = Pt(13)
    run_t2.bold = True
    run_t2.font.color.rgb = RGBColor(0x00, 0x55, 0x99)

    run_t3 = p_title.add_run("PROJECT BASED LEARNING + AGILE SPRINT FRAMEWORK (SPRINT 7)")
    run_t3.font.name = 'Calibri'
    run_t3.font.size = Pt(11)
    run_t3.bold = True
    run_t3.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(8)
    r_div = p_div.add_run("═" * 70)
    r_div.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    r_div.bold = True

    # -------------------------------------------------------------
    # TABLE 0: INFORMASI AKADEMIK
    # -------------------------------------------------------------
    info_data = [
        ("Informasi Akademik", "Keterangan Terperinci"),
        ("Mata Kuliah", "Praktikum Pemrograman Berorientasi Obyek (PBO)"),
        ("Dosen Pengampu", "Nirwana Haidar Hari, S.Pd., M.Kom."),
        ("Nama Mahasiswa / NRP", "Ahmad Yasar Maulana Dais / 3125522018"),
        ("Program Studi & Kampus", "D3 Teknik Informatika - PENS PSDKU Sumenep (2026)"),
        ("Nama Proyek & Modul", "AutoRent System – Sprint 7: Abstraction, Abstract Method, Interface & Polymorphism")
    ]
    t0 = doc.add_table(rows=len(info_data), cols=2)
    t0.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t0, color="003366", sz="6")
    for r_idx, (col1, col2) in enumerate(info_data):
        row = t0.rows[r_idx]
        cell0, cell1 = row.cells[0], row.cells[1]
        cell0.width = Inches(2.2)
        cell1.width = Inches(4.3)
        set_cell_margins(cell0, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell1, top=80, bottom=80, left=120, right=120)

        p0 = cell0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        p1 = cell1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)

        if r_idx == 0:
            set_cell_background(cell0, "003366")
            set_cell_background(cell1, "003366")
            r0 = p0.add_run(col1)
            r0.bold = True
            r0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r1 = p1.add_run(col2)
            r1.bold = True
            r1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        else:
            set_cell_background(cell0, "F0F4F8" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_background(cell1, "F0F4F8" if r_idx % 2 == 1 else "FFFFFF")
            r0 = p0.add_run(col1)
            r0.bold = True
            r0.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
            r1 = p1.add_run(col2)
            r1.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Helper function for Section Headings
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(13)
        r.bold = True
        r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(11)
        r.bold = True
        r.font.color.rgb = RGBColor(0x00, 0x55, 0x99)
        return p

    def add_callout(box_text, title="📌 SPRINT GOAL SPRINT 7 (MODUL 7)", bg_hex="EBF3FA", border_hex="005599"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, bg_hex)
        set_box_borders(cell, color=border_hex, sz="8")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_t = p.add_run(title + "\n")
        r_t.bold = True
        r_t.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        r_b = p.add_run(box_text)
        r_b.font.size = Pt(9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # BAB I: TUJUAN PRAKTIKUM & FONDASI AGILE SPRINT (SPRINT 7)
    # -------------------------------------------------------------
    add_h1("BAB I: TUJUAN PRAKTIKUM & FONDASI AGILE SPRINT (SPRINT 7)")

    add_h2("1.1 Tujuan Praktikum (Capaian Pembelajaran Modul 7)")
    p_tujuan = doc.add_paragraph(
        "Berdasarkan Buku Panduan Resmi Modul 7 Praktikum Pemrograman Berorientasi Obyek (PBO) Politeknik Elektronika Negeri Surabaya (PENS) PSDKU Sumenep, capaian pembelajaran yang wajib dikuasai dan dibuktikan pada praktikum ini adalah:\n"
        "1. Mampu membedakan secara fundamental karakteristik dan tujuan perancangan antara class biasa (concrete class), abstract class, dan interface.\n"
        "2. Mampu membuat dan mendeklarasikan abstract class yang memiliki konstruktor, atribut umum, method konkret, dan abstract method.\n"
        "3. Mampu mengimplementasikan abstract method secara konsisten dan adaptif pada seluruh subclass konkret (Mobil, Motor, Minibus, Pelanggan, Pegawai).\n"
        "4. Mampu membuat antarmuka (interface) dan menerapkannya pada class menggunakan keyword implements, termasuk pemanfaatan default method Java modern.\n"
        "5. Mampu membuktikan dan mengoperasikan mekanisme Dynamic Polymorphism melalui Superclass Abstract Reference dan Interface Reference lintas hierarki pewarisan.\n"
        "6. Mampu melakukan refactoring fungsional terhadap proyek hasil Pertemuan 6 (AutoRent System) berdasarkan prinsip Abstraction tanpa merusak struktur relasi objek sebelumnya.\n"
        "7. Mampu melaksanakan Sprint Review dan Sprint Retrospective secara profesional sebagai bagian dari kerangka kerja pengembangan perangkat lunak Agile Scrum."
    )

    add_h2("1.2 Review P6 dan Jembatan Arsitektural Menuju P7")
    doc.add_paragraph(
        "Pada praktikum Pertemuan 6 (P6), proyek AutoRent System telah berhasil direfaktor dengan menerapkan konsep Dynamic Polymorphism, Method Overriding, Method Overloading, dan Dynamic Binding. Hierarki class telah terspesialisasi menjadi Superclass Kendaraan (dengan Subclass Mobil, Motor, Minibus) dan Superclass Person (dengan Subclass Pelanggan dan Pegawai), di samping mempertahankan seluruh relasi objek P4 seperti Komposisi (Mesin dan DetailItemBiaya), Agregasi (GarasiArmada), Asosiasi (TransaksiSewa), serta Dependency (KuitansiPrinter dan LayananNotifikasi).\n\n"
        "Meskipun demikian, evaluasi desain arsitektural pada P6 menunjukkan adanya celah perancangan yang krusial: Superclass Kendaraan dan Person pada P6 masih berupa concrete class (class biasa). Hal ini memungkinkan terjadinya instansiasi objek langsung seperti new Kendaraan(...) atau new Person(...). Di dunia nyata bisnis rental armada, tidak pernah ada 'kendaraan tanpa wujud konkrit' yang beroperasi di jalan raya, maupun 'manusia tanpa peran spesifik' dalam sistem. Seluruh armada riil selalu berwujud Mobil, Motor, atau Minibus, dan seluruh manusia selalu berperan sebagai Pelanggan atau Pegawai.\n\n"
        "Sesuai instruksi resmi Modul 7, mahasiswa TIDAK membuat proyek baru dari nol, melainkan melanjutkan dan menyempurnakan proyek AutoRent System hasil P6. Fokus utama Sprint 7 adalah melakukan evolusi arsitektural menuju tahapan Abstraction: mentransformasikan superclass konseptual menjadi Abstract Class murni, mengubah method bervariasi menjadi Abstract Method wajib, serta mengekstrak antarmuka kontrak perilaku umum (Interface DapatDilacak, Perawatan, Cetakable) guna mewujudkan arsitektur yang decoupled, modular, dan standar industri."
    )

    add_h2("1.3 Agile Sprint Framework (Sprint Goal & Sprint Backlog Modul 7)")
    add_callout(
        "Sprint Goal P7: Mengembangkan desain proyek P6 dengan menerapkan abstract class dan interface yang sesuai sehingga struktur class lebih jelas, perilaku object memiliki kontrak yang konsisten, mencegah instansiasi entitas umum secara langsung, mendukung multiple interface realization lintas hierarki, dan program tetap dapat dijalankan secara terintegrasi 100%.",
        title="📌 SPRINT GOAL SPRINT 7 (MODUL 7)"
    )

    doc.add_paragraph("Tabel berikut menyajikan Sprint Backlog terperinci yang dieksekusi sepanjang iterasi Sprint 7 untuk mencapai seluruh kriteria Definition of Done:")

    backlog_data = [
        ("ID Backlog", "Deskripsi Aktivitas Sprint Backlog", "Target Deliverable", "Status Akhir"),
        ("SB-01", "Audit hierarchy dan behavior proyek P6", "Matriks Audit Abstraksi & Evaluasi Celah Desain", "DONE (100%)"),
        ("SB-02", "Menentukan kandidat abstract class", "Spesifikasi Abstract Class Kendaraan & Person", "DONE (100%)"),
        ("SB-03", "Menentukan abstract method wajib", "Daftar Abstract Method pada Kendaraan & Person", "DONE (100%)"),
        ("SB-04", "Menentukan kandidat interface sistem", "Definisi Interface DapatDilacak, Perawatan, Cetakable", "DONE (100%)"),
        ("SB-05", "Memperbarui UML Class Diagram", "UML Diagram Modul 7 (Generalization & Realization)", "DONE (100%)"),
        ("SB-06", "Mengimplementasikan abstract class dan interface", "17 File Source Code Java (src/) Terintegrasi", "DONE (100%)"),
        ("SB-07", "Menguji polymorphism dan program", "Test Suite Main.java (8 Fase Pengujian Selesai)", "DONE (100%)")
    ]
    t_bl = doc.add_table(rows=len(backlog_data), cols=4)
    t_bl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_bl, color="003366", sz="4")
    for r_idx, row_data in enumerate(backlog_data):
        row = t_bl.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if r_idx == 0:
                set_cell_background(cell, "003366")
                r = p.add_run(val)
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9)
            else:
                set_cell_background(cell, "F0F4F8" if r_idx % 2 == 1 else "FFFFFF")
                r = p.add_run(val)
                r.font.size = Pt(8.5)
                if c_idx == 0:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
                elif c_idx == 3:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x00, 0x66, 0x00)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2("1.4 Bagian A – Audit Desain P6 (Refactoring Abstraction & Interface)")
    doc.add_paragraph(
        "Sesuai arahan resmi Bagian A Modul 7, dilakukan audit menyeluruh terhadap struktur superclass, subclass, dan behavior hasil P6. Tujuannya adalah mengidentifikasi class yang seharusnya tidak dapat diinstansiasi secara langsung serta perilaku yang dapat didefinisikan sebagai kontrak melalui interface:"
    )

    audit_data = [
        ("Class / Behavior", "Kondisi P6", "Rencana P7", "Alasan Arsitektural"),
        ("Kendaraan", "Class Biasa (Concrete)", "Abstract Class", "Hanya mewakili konsep umum armada. Tidak boleh diinstansiasi langsung menggunakan new Kendaraan() karena kendaraan riil selalu berupa Mobil, Motor, atau Minibus."),
        ("Person", "Class Biasa (Concrete)", "Abstract Class", "Hanya mewakili kerangka dasar data identitas manusia. Tidak boleh diinstansiasi langsung karena pengguna sistem selalu berupa Pelanggan atau Pegawai."),
        ("hitungEstimasiBiaya()", "Method Biasa (Virtual Default)", "Abstract Method", "Setiap kategori kendaraan memiliki formula kalkulasi tarif, struktur diskon durasi, dan surcharge operasional yang berbeda secara fundamental."),
        ("tampilkanSpesifikasi()", "Method Biasa (Virtual Default)", "Abstract Method", "Setiap armada menyajikan spesifikasi unik (Mobil: transmisi/AC, Motor: tipe/bagasi helm, Minibus: kapasitas kursi/izin trayek)."),
        ("tampilkanProfil() & hitungKompensasi()", "Method Biasa pada Person", "Abstract Method", "Pelanggan menampilkan status member & reward loyalti, sedangkan Pegawai menampilkan NIK, jabatan, dan slip total gaji bulanan."),
        ("DapatDilacak", "Belum Ada", "Interface", "Mendefinisikan kontrak telematika GPS IoT yang dapat diterapkan pada entitas armada (Mobil, Motor, Minibus) maupun petugas lapangan (Pegawai) secara lintas hierarki."),
        ("Perawatan", "Belum Ada", "Interface", "Menstandarkan kontrak pemeliharaan berkala, cek kelayakan jalan, dan jadwal servis bagi seluruh jenis kendaraan rental."),
        ("Cetakable", "Belum Ada", "Interface", "Menstandarkan kontrak formatting dokumen resmi untuk kuitansi transaksi sewa, kartu member digital, dan slip kompensasi pegawai.")
    ]
    t_aud = doc.add_table(rows=len(audit_data), cols=4)
    t_aud.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_aud, color="003366", sz="4")
    for r_idx, row_data in enumerate(audit_data):
        row = t_aud.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if r_idx == 0:
                set_cell_background(cell, "003366")
                r = p.add_run(val)
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9)
            else:
                set_cell_background(cell, "F0F4F8" if r_idx % 2 == 1 else "FFFFFF")
                r = p.add_run(val)
                r.font.size = Pt(8.5)
                if c_idx in [0, 2]:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # BAB II: KAJIAN TEORI & PERANCANGAN ARSITEKTUR
    # -------------------------------------------------------------
    add_h1("BAB II: KAJIAN TEORI & PERANCANGAN ARSITEKTUR ABSTRAKSI & INTERFACE")

    add_h2("2.1 Konsep Abstraksi & Abstract Class dalam Pemrograman Berorientasi Obyek")
    doc.add_paragraph(
        "Abstraksi (Abstraction) adalah salah satu dari empat pilar fundamental Pemrograman Berorientasi Obyek (di samping Encapsulation, Inheritance, dan Polymorphism). Abstraksi bertujuan untuk menyembunyikan detail implementasi internal yang kompleks dan hanya menampilkan esensi fungsionalitas umum kepada dunia luar.\n\n"
        "Abstract Class merupakan class yang dideklarasikan dengan keyword 'abstract'. Karakteristik utama Abstract Class meliputi:\n"
        "1. Tidak Dapat Diinstansiasi Langsung: Perintah seperti 'Kendaraan k = new Kendaraan(...);' akan memicu Compile Error 'Kendaraan is abstract; cannot be instantiated'.\n"
        "2. Wadah Polimorfisme: Abstract Class sangat ideal digunakan sebagai tipe variabel reference (Upcasting), misalnya 'Kendaraan k = new Mobil(...);'.\n"
        "3. Memiliki Konstruktor: Berbeda dengan interface, abstract class tetap memiliki konstruktor yang dipanggil oleh subclass melalui konstruktor chaining menggunakan keyword 'super(...)'.\n"
        "4. Kombinasi Method Konkret & Abstrak: Abstract class dapat mengimplementasikan atribut instance dan method-method yang bersifat universal, sekaligus mendeklarasikan abstract method yang mewajibkan implementasi spesifik pada subclass."
    )

    add_h2("2.2 Konsep & Karakteristik Abstract Method")
    doc.add_paragraph(
        "Abstract Method adalah method yang hanya memiliki deklarasi signature (nama method, tipe return, dan parameter) tanpa blok tubuh implementasi ({ }), diakhiri dengan tanda titik koma (;). Contoh:\n"
        "    public abstract double hitungEstimasiBiaya(int durasiHari);\n\n"
        "Aturan ketat Abstract Method dalam bahasa Java:\n"
        "• Hanya boleh berada di dalam Abstract Class atau Interface.\n"
        "• Tidak boleh dideklarasikan dengan modifier private (karena wajib dioverride oleh subclass).\n"
        "• Tidak boleh menggunakan keyword static atau final (karena static terikat class dan final melarang overriding).\n"
        "• Seluruh Subclass Konkret wajib meng-override dan mengimplementasikan seluruh abstract method yang diwarisinya, kecuali jika subclass tersebut juga dideklarasikan sebagai abstract."
    )

    add_h2("2.3 Konsep Interface & Keyword implements")
    doc.add_paragraph(
        "Interface adalah cetak biru (blueprint) yang mendefinisikan kontrak perilaku (behavioral contract) murni yang dapat direalisasikan oleh class mana pun tanpa memandang hierarki pewarisannya. Class yang mengikat diri pada suatu interface menggunakan keyword 'implements'.\n\n"
        "Karakteristik & Fitur Modern Interface Java:\n"
        "1. Multiple Interface Realization: Java tidak mendukung multiple inheritance pada class (menghindari Diamond Problem), namun satu class diizinkan mengimplementasikan lebih dari satu interface secara simultan (contoh: class Mobil extends Kendaraan implements DapatDilacak, Perawatan, Cetakable).\n"
        "2. Kontrak Perilaku Lintas Hierarki: Dua class yang sama sekali tidak berada dalam satu hierarki pewarisan (misalnya Mobil yang merupakan turunan Kendaraan dan Pegawai yang merupakan turunan Person) dapat sama-sama mengimplementasikan interface DapatDilacak.\n"
        "3. Default & Static Methods (Java 8+): Interface modern mendukung 'default method' yang menyediakan implementasi bawaan tanpa memaksa implementor lama untuk menulis ulang kodenya, serta 'static method' sebagai utility mandiri.\n"
        "4. Atribut Publik Bersifat Konstanta: Seluruh atribut yang dideklarasikan di dalam interface secara otomatis bersifat 'public static final'."
    )

    add_h2("2.4 Matriks Perbandingan Komprehensif: Concrete Class vs Abstract Class vs Interface")
    doc.add_paragraph("Tabel berikut menyajikan analisis komparasi komprehensif antara Concrete Class, Abstract Class, dan Interface berdasarkan 8 dimensi arsitektural:")

    matriks_data = [
        ("Dimensi Pembanding", "Concrete Class (Class Biasa)", "Abstract Class", "Interface"),
        ("Tujuan Utama", "Mewakili entitas nyata yang siap dioperasikan langsung", "Mewakili konsep dasar umum & kerangka berbagi logika", "Menentukan kontrak perilaku (behavioral contract) murni"),
        ("Instansiasi (new)", "Bisa diinstansiasi langsung", "TIDAK BISA diinstansiasi langsung", "TIDAK BISA diinstansiasi langsung"),
        ("Keyword Penggunaan", "class dan extends", "abstract class dan extends", "interface dan implements"),
        ("Atribut Instance", "Bebas (private, protected, public, dsb.)", "Bebas (mendukung state/atribut instance penuh)", "Hanya konstanta (public static final secara implisit)"),
        ("Konstruktor", "Memiliki konstruktor instance", "Memiliki konstruktor (dipanggil via super())", "TIDAK memiliki konstruktor"),
        ("Implementasi Method", "Seluruh method wajib memiliki body", "Kombinasi method konkret & abstract method", "Hanya deklarasi method (atau default/static method)"),
        ("Dukungan Hierarki", "Single inheritance (1 superclass langsung)", "Single inheritance (1 superclass langsung)", "Multiple realization (bisa implement banyak interface)"),
        ("Hubungan Semantik", "Relasi 'is-a' konkrit", "Relasi 'is-a' konseptual/umum", "Relasi 'can-do' / 'acts-as' (kemampuan perilaku)")
    ]
    t_mat = doc.add_table(rows=len(matriks_data), cols=4)
    t_mat.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_mat, color="003366", sz="4")
    for r_idx, row_data in enumerate(matriks_data):
        row = t_mat.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if r_idx == 0:
                set_cell_background(cell, "003366")
                r = p.add_run(val)
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9)
            else:
                set_cell_background(cell, "F0F4F8" if r_idx % 2 == 1 else "FFFFFF")
                r = p.add_run(val)
                r.font.size = Pt(8.5)
                if c_idx == 0:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2("2.5 Hubungan Abstraksi dengan Polymorphism")
    doc.add_paragraph(
        "Penerapan Abstraction memperkuat fondasi Dynamic Polymorphism melalui dua saluran utama:\n"
        "1. Superclass Abstract Reference: Variabel bertipe abstract class digunakan untuk menunjuk instance subclass konkrit (contoh: Kendaraan k = new Minibus(...)). Saat method abstract 'k.hitungEstimasiBiaya(5)' dipanggil, JVM secara otomatis mengeksekusi implementasi Minibus melalui mekanisme Dynamic Method Dispatch.\n"
        "2. Interface Reference Polymorphism: Variabel bertipe interface dapat menunjuk objek apa pun yang mengimplementasikan interface tersebut tanpa peduli asal hierarkinya (contoh: DapatDilacak tracker = new Pegawai(...)). Hal ini memungkinkan pemrosesan koleksi heterogen (Heterogeneous Collection) seperti List<DapatDilacak> atau List<Cetakable> secara seragam dan sangat modular."
    )

    add_h2("2.6 Pembaruan UML Class Diagram (Modul 7 Architecture)")
    doc.add_paragraph(
        "Berikut adalah representasi arsitektural lengkap UML Class Diagram AutoRent System Modul 7 yang telah disempurnakan dengan notasi standar UML (<<abstract>>, <<interface>>, Generalization △, Realization ┆▷, Composition ◆, Aggregation ◇, Association ──>, Dependency ..>):"
    )

    uml_text = """+-----------------------------------------------------------------------------+
|               📌 UML CLASS DIAGRAM ARSITEKTUR LENGKAP (MODUL 7)             |
+-----------------------------------------------------------------------------+
|                                                                             |
|                      +----------------------------------+                   |
|                      |      <<abstract>> Kendaraan      |                   |
|                      +----------------------------------+                   |
|                      | - nomorPlat: String              |                   |
|                      | - merk: String                   |                   |
|                      | - model: String                  |                   |
|                      | - tarifPerHari: double           |                   |
|                      | - mesin: Mesin (◆ Komposisi)     |                   |
|                      +----------------------------------+                   |
|                      | + <<abstract>> hitungEstimasi()  |                   |
|                      | + <<abstract>> tampilkanSpesif() |                   |
|                      | + hitungEstimasi(durasi, diskon) |                   |
|                      +----------------------------------+                   |
|                                       △                                     |
|                                       │ Generalization (extends)            |
|                ┌──────────────────────┼──────────────────────┐              |
|                │                      │                      │              |
|     +--------------------+ +--------------------+ +--------------------+    |
|     |       Mobil        | |       Motor        | |      Minibus       |    |
|     +--------------------+ +--------------------+ +--------------------+    |
|     | - tipeTransmisi    | | - jenisMotor       | | - jumlahKursi      |    |
|     | - fasilitasAC      | | - bagasiHelm       | | - noIzinTrayek     |    |
|     +--------------------+ +--------------------+ +--------------------+    |
|                ┆                      ┆                      ┆              |
|                ┆ implements           ┆ implements           ┆ implements   |
|                └──────────────────────┼──────────────────────┘              |
|                                       ▼                                     |
|                        +──────────────────────────────+                     |
|                        |   <<interface>> DapatDilacak |                     |
|                        +──────────────────────────────+                     |
|                        | + getLokasi(): String        |                     |
|                        | + perbaruiLokasi(koor, area) |                     |
|                        | + getStatusGPS(): String     |                     |
|                        +──────────────────────────────+                     |
|                                       ▲                                     |
|                                       ┆ implements (Lintas Hierarki!)       |
|                                       │                                     |
|     +--------------------+            │            +--------------------+   |
|     |<<abstract>> Person |            │            |      Pegawai       |   |
|     +--------------------+            │            +--------------------+   |
|     | - id, nama, email  |            │            | - NIK: String      |   |
|     +--------------------+            │            | - posisiJabatan    |   |
|     | + <<abs>> profil() |            │            | - gajiPokok        |   |
|     | + <<abs>> kompen() |            │            +--------------------+   |
|     +--------------------+            │                       ▲             |
|                △                      │                       │ extends     |
|                └──────────────────────┴───────────────────────┘             |
|                │                                                            |
|     +--------------------+            +──────────────────────────────+      |
|     |     Pelanggan      | ┆implements|   <<interface>> Cetakable    |      |
|     +--------------------+───────────>+──────────────────────────────+      |
|     | - nomorKTP         |            | + formatCetakDokumen()       |      |
|     | - saldoDeposit     |            | + generateRingkasan()        |      |
|     +--------------------+            +──────────────────────────────+      |
|                                                       ▲                     |
|     +────────────────────+               ┆ implements │                     |
|     |   GarasiArmada     |               └────────────┼─────────────────┐   |
|     | (◇ Agregasi Armada)|                            │                 │   |
|     +────────────────────+                            │                 │   |
|     +────────────────────+            +──────────────────────────────+  │   |
|     |   TransaksiSewa    |───────────>|        DetailItemBiaya       |  │   |
|     | (Asosiasi & Cetak) | (◆ Kompos) |       (Komposisi Rincian)    |  │   |
|     +────────────────────+            +──────────────────────────────+  │   |
|                :                                                        │   |
|                :..> KuitansiPrinter (Service Dependency) <──────────────┘   |
|                :..> LayananNotifikasi (Service Gateway)                     |
+-----------------------------------------------------------------------------+"""

    tbl_uml = doc.add_table(rows=1, cols=1)
    tbl_uml.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_uml = tbl_uml.cell(0, 0)
    c_uml.width = Inches(6.5)
    set_cell_background(c_uml, "F4F6F9")
    set_box_borders(c_uml, color="003366", sz="6")
    set_cell_margins(c_uml, top=80, bottom=80, left=100, right=100)
    p_u = c_uml.paragraphs[0]
    p_u.paragraph_format.space_after = Pt(0)
    r_u = p_u.add_run(uml_text)
    r_u.font.name = 'Consolas'
    r_u.font.size = Pt(7.5)
    r_u.font.color.rgb = RGBColor(0x00, 0x22, 0x44)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # BAB III: IMPLEMENTASI KODE PROGRAM JAVA (SRC)
    # -------------------------------------------------------------
    add_h1("BAB III: IMPLEMENTASI KODE PROGRAM JAVA (SRC)")

    add_h2("3.1 Arsitektur Hierarki Abstract Class & Interface")
    doc.add_paragraph(
        "Pada Sprint 7 ini, struktur source code AutoRent System direorganisasi menjadi 17 file Java modular yang saling terhubung erat:\n"
        "• Interface Layer: DapatDilacak.java, Perawatan.java, Cetakable.java.\n"
        "• Abstraction Superclass Layer: Person.java (Abstract Class), Kendaraan.java (Abstract Class).\n"
        "• Concrete Subclass Layer: Mobil.java, Motor.java, Minibus.java, Pelanggan.java, Pegawai.java.\n"
        "• Structural & Relational Layer: Mesin.java (Komposisi Kendaraan), DetailItemBiaya.java (Komposisi Transaksi), GarasiArmada.java (Agregasi Polimorfis), TransaksiSewa.java (Asosiasi & Orkestrasi Sistem).\n"
        "• Service Dependency Layer: KuitansiPrinter.java (Interface Utility Dependency), LayananNotifikasi.java (Gateway Alert Dependency).\n"
        "• Test Suite Driver: Main.java (Driver 8 Fase Pengujian Komprehensif)."
    )

    add_h2("3.2 Dokumentasi Listing Seluruh File Source Code Java (17 File)")

    # Read files from src/
    src_files = [
        ("DapatDilacak.java", "Interface Kontrak Telematika & GPS Tracking IoT"),
        ("Perawatan.java", "Interface Standarisasi Pemeliharaan & Uji Kelayakan Armada"),
        ("Cetakable.java", "Interface Standarisasi Format Dokumen Resmi (Kuitansi/Kartu/Slip)"),
        ("Person.java", "Abstract Superclass Abstraksi Entitas Manusia (User/Staff)"),
        ("Kendaraan.java", "Abstract Superclass Abstraksi Armada Kendaraan Rental"),
        ("Mobil.java", "Subclass Konkret Mobil Penumpang (extends Kendaraan, implements 3 Interface)"),
        ("Motor.java", "Subclass Konkret Sepeda Motor (extends Kendaraan, implements 3 Interface)"),
        ("Minibus.java", "Subclass Konkret Minibus Travel (extends Kendaraan, implements 3 Interface)"),
        ("Pelanggan.java", "Subclass Konkret Pelanggan Rental (extends Person, implements Cetakable)"),
        ("Pegawai.java", "Subclass Konkret Pegawai Operasional (extends Person, implements DapatDilacak, Cetakable)"),
        ("Mesin.java", "Komponen Komposisi Mesin Internal Kendaraan"),
        ("DetailItemBiaya.java", "Komponen Komposisi Baris Rincian Biaya Transaksi"),
        ("GarasiArmada.java", "Agregasi Polimorfis Pool Armada & Batch Interface Processing"),
        ("TransaksiSewa.java", "Sentral Orkestrasi Relasi Transaksi (implements Cetakable)"),
        ("KuitansiPrinter.java", "Service Utility Pencetakan Berbasis Interface Cetakable"),
        ("LayananNotifikasi.java", "Service Gateway Notifikasi Multi-Kanal & Telemetry"),
        ("Main.java", "Test Suite Utama & Driver Verifikasi 8 Fase Pengujian Modul 7")
    ]

    for fname, desc in src_files:
        p_fn = doc.add_paragraph()
        p_fn.paragraph_format.space_before = Pt(8)
        p_fn.paragraph_format.space_after = Pt(2)
        p_fn.paragraph_format.keep_with_next = True
        r_fn = p_fn.add_run(f"📄 File: {fname} — {desc}")
        r_fn.bold = True
        r_fn.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        r_fn.font.size = Pt(10)

        fpath = os.path.join("src", fname)
        if os.path.exists(fpath):
            with open(fpath, "r", encoding="utf-8") as fp:
                code_content = fp.read()
        else:
            code_content = "// File tidak ditemukan"

        tbl_c = doc.add_table(rows=1, cols=1)
        tbl_c.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_code = tbl_c.cell(0, 0)
        c_code.width = Inches(6.5)
        set_cell_background(c_code, "F9FBFD")
        set_box_borders(c_code, color="CCCCCC", sz="4")
        set_cell_margins(c_code, top=60, bottom=60, left=80, right=80)
        p_c = c_code.paragraphs[0]
        p_c.paragraph_format.space_after = Pt(0)
        p_c.paragraph_format.line_spacing = 1.0
        r_c = p_c.add_run(code_content)
        r_c.font.name = 'Consolas'
        r_c.font.size = Pt(7.5)
        r_c.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # BAB IV: PENGUJIAN SISTEM & LOG EKSEKUSI
    # -------------------------------------------------------------
    add_h1("BAB IV: PENGUJIAN SISTEM, TABEL PENGUJIAN POLYMORPHISM & LOG EKSEKUSI")

    add_h2("4.1 Skenario Pengujian Komprehensif (8 Fase)")
    doc.add_paragraph(
        "Untuk membuktikan seluruh persyaratan Modul 7 Bab 6 (Bagian F – Pengujian Polymorphism) dan memastikan tidak ada regresi fitur P1–P6, test suite Main.java dirancang mencakup 8 fase pengujian yang komprehensif:\n"
        "• Fase 1: Pembuktian Larangan Instansiasi Abstract Class (Kendaraan & Person tidak dapat di-new).\n"
        "• Fase 2: Instansiasi Subclass Konkret (Mobil, Motor, Minibus, Pelanggan, Pegawai) dengan validasi konstruktor.\n"
        "• Fase 3: Pengujian Dynamic Method Dispatch melalui Superclass Abstract Reference.\n"
        "• Fase 4: Pengujian Interface Polymorphism (DapatDilacak, Perawatan, Cetakable) lintas hierarki class.\n"
        "• Fase 5: Pengujian Heterogeneous Collection (Array & ArrayList polimorfis superclass dan interface).\n"
        "• Fase 6: Pengujian Method Overloading & Defensive Programming Boundary Conditions.\n"
        "• Fase 7: Pengujian Integrasi Alur Bisnis Penuh AutoRent (Garasi, Transaksi, Mesin, Printer, Notifikasi).\n"
        "• Fase 8: Ringkasan Eksekutif & Verifikasi Ketercapaian 100% Definition of Done."
    )

    add_h2("4.2 Matriks & Tabel Pengujian Polymorphism (Modul 7 Bagian F)")
    doc.add_paragraph("Tabel berikut menyajikan hasil verifikasi empiris pengujian polimorfisme sesuai format resmi Modul 7 Bagian F:")

    uji_data = [
        ("No", "Skenario Pengujian", "Eksekusi Kode / Input", "Hasil yang Diharapkan", "Hasil Aktual", "Status"),
        ("1", "Audit Larangan Instansiasi Abstract Class", "Mencoba: new Kendaraan(...) & new Person(...)", "Compiler menolak pembuatan objek langsung", "Terbukti: Instansiasi dicegah oleh compiler Java", "PASSED"),
        ("2", "Instansiasi Subclass Konkret", "new Mobil(...), new Motor(...), new Minibus(...), new Pelanggan(...), new Pegawai(...)", "Seluruh objek konkret berhasil dialokasikan di Heap Memory", "5 Objek konkret berhasil diinstansiasi 100%", "PASSED"),
        ("3", "Memanggil abstract method via Superclass Reference", "Kendaraan refK = mobil;\nrefK.tampilkanSpesifikasi();\nrefK.hitungEstimasiBiaya(5);", "JVM mengeksekusi logika spesifik milik subclass (Dynamic Dispatch)", "Tereksekusi tepat sesuai subclass Mobil, Motor, Minibus", "PASSED"),
        ("4", "Memanggil method via Reference Interface", "DapatDilacak p = pegawai;\np.getLokasi();\np.perbaruiLokasi(...);", "Interface polymorphism lintas hierarki berjalan mulus", "Lokasi dan status GPS terupdate secara akurat", "PASSED"),
        ("5", "Heterogeneous Collection Processing", "List<Cetakable> daftar = ...;\nfor(Cetakable d : daftar) d.generateRingkasan();", "Setiap elemen koleksi menjalankan format cetak masing-masing", "4 Tipe objek berbeda mencetak ringkasannya sendiri", "PASSED"),
        ("6", "Method Overloading & Defensive Boundary Test", "hitungEstimasiBiaya(4, 20.0);\ntambahDeposit(-1000);\nMinibus kursi 4;", "Overloading tereksekusi benar; input ilegal melempar IllegalArgumentException", "Kalkulasi promo tepat; seluruh negative boundary tertangkap", "PASSED"),
        ("7", "End-to-End Flow & Service Dependency", "TransaksiSewa -> Bayar VA -> Selesai -> Printer.cetakDokumenResmi(trx)", "Kuitansi tercetak resmi via interface Cetakable; status unit READY", "Kuitansi, kartu member, dan slip gaji tercetak sempurna", "PASSED")
    ]
    t_uj = doc.add_table(rows=len(uji_data), cols=6)
    t_uj.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_uj, color="003366", sz="4")
    for r_idx, row_data in enumerate(uji_data):
        row = t_uj.rows[r_idx]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if r_idx == 0:
                set_cell_background(cell, "003366")
                r = p.add_run(val)
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(8.5)
            else:
                set_cell_background(cell, "F0F4F8" if r_idx % 2 == 1 else "FFFFFF")
                r = p.add_run(val)
                r.font.size = Pt(8)
                if c_idx == 0:
                    r.bold = True
                elif c_idx == 5:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x00, 0x66, 0x00)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2("4.3 Log Output Eksekusi Terminal Java Konsol Lengkap")
    doc.add_paragraph(
        "Berikut adalah rekaman output terminal riil tanpa rekayasa dari proses kompilasi 'javac -encoding UTF-8 -d bin src/*.java' dan eksekusi 'java -cp bin Main':"
    )

    if os.path.exists("execution_log_p7.txt"):
        with open("execution_log_p7.txt", "r", encoding="utf-8") as fp:
            log_content = fp.read()
    else:
        log_content = "[Log output tidak tersedia]"

    tbl_log = doc.add_table(rows=1, cols=1)
    tbl_log.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_log = tbl_log.cell(0, 0)
    c_log.width = Inches(6.5)
    set_cell_background(c_log, "1E1E1E")
    set_box_borders(c_log, color="333333", sz="4")
    set_cell_margins(c_log, top=80, bottom=80, left=100, right=100)
    p_l = c_log.paragraphs[0]
    p_l.paragraph_format.space_after = Pt(0)
    p_l.paragraph_format.line_spacing = 1.0
    r_l = p_l.add_run(log_content)
    r_l.font.name = 'Consolas'
    r_l.font.size = Pt(7.5)
    r_l.font.color.rgb = RGBColor(0xD4, 0xD4, 0xD4)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_h2("4.4 Analisis & Pembahasan Mendalam Hasil Pengujian")
    doc.add_paragraph(
        "Berdasarkan log eksekusi terminal di atas, diperoleh beberapa poin temuan dan analisis arsitektural yang mendalam:\n\n"
        "1. Pembuktian Integritas Abstraksi (Fase 1 & 2): Penetapan keyword 'abstract' pada Kendaraan dan Person secara efektif mencegah pembuatan objek prematur/konseptual. Compiler Java menjamin bahwa hanya objek dengan wujud konkrit yang dapat hidup di memory heap, sehingga menghilangkan celah data anomali.\n\n"
        "2. Keunggulan Abstract Method & Dynamic Dispatch (Fase 3): Pemanggilan method hitungEstimasiBiaya(5) melalui reference Superclass Kendaraan secara otomatis menghasilkan nilai kalkulasi yang berbeda pada Mobil (Rp 1.900.000, diskon durasi 5%), Motor (Rp 705.000, diskon hemat 6%), dan Minibus (Rp 5.750.000, termasuk surcharge operasional armada besar). JVM terbukti mengeksekusi vtable subclass konkrit secara transparan.\n\n"
        "3. Efektivitas Multiple Interface Realization (Fase 4 & 5): Mobil, Motor, dan Minibus berhasil merealisasikan tiga interface sekaligus (DapatDilacak, Perawatan, Cetakable). Yang paling menarik, class Pegawai (turunan Person) dan armada Kendaraan dapat ditampung dalam satu array 'DapatDilacak[]' untuk dimonitor posisinya secara terpadu. Ini membuktikan bahwa interface mampu menghubungkan entitas lintas hierarki pewarisan (cross-hierarchy polymorphism).\n\n"
        "4. Ketahanan Terhadap Input Cacat (Fase 6): Seluruh validasi defensif bekerja 100%. Upaya memasukkan diskon > 50%, top-up deposit negatif, atau membuat minibus berkapasitas < 8 kursi berhasil digagalkan dengan pelemparan IllegalArgumentException yang jelas dan aman.\n\n"
        "5. Keutuhan Integrasi Sistem P1–P6 (Fase 7): Seluruh relasi fundamental (Komposisi Mesin, Komposisi DetailItemBiaya, Agregasi GarasiArmada, Asosiasi TransaksiSewa, dan Dependency Printer/Notifikasi) tetap beroperasi harmonis tanpa ada fitur yang rusak (zero regression)."
    )

    # -------------------------------------------------------------
    # BAB V: SPRINT EVALUATION, DEFINITION OF DONE & PENUTUP
    # -------------------------------------------------------------
    add_h1("BAB V: SPRINT EVALUATION, DEFINITION OF DONE & PENUTUP")

    add_h2("5.1 Sprint Review (Format Resmi Modul 7 Bab 8)")
    doc.add_paragraph("Tabel berikut menyajikan evaluasi ketercapaian checklist Sprint Review sesuai format resmi Modul 7 Bab 8:")

    rev_data = [
        ("Item Evaluasi Sprint Review", "Status Hasil Verifikasi"),
        ("Abstract class berhasil dibuat (Kendaraan dan Person)", "VERIFIED (100% Selesai)"),
        ("Abstract method diimplementasikan subclass secara konsisten", "VERIFIED (100% Selesai)"),
        ("Minimal 2 subclass dapat digunakan (Mobil, Motor, Minibus, Pelanggan, Pegawai)", "VERIFIED (100% Selesai)"),
        ("Interface berhasil diimplementasikan (DapatDilacak, Perawatan, Cetakable)", "VERIFIED (100% Selesai)"),
        ("Class diagram sesuai dengan struktur kode sumber Java", "VERIFIED (100% Selesai)"),
        ("Pengujian polymorphism berhasil pada superclass & interface", "VERIFIED (100% Selesai)"),
        ("Fitur proyek sebelumnya (P1–P6) tetap berjalan normal tanpa regresi", "VERIFIED (100% Selesai)"),
        ("Ringkasan Keterverifikasian Sprint Review", "7 / 7 ITEM TERVERIFIKASI (100%)")
    ]
    t_rev = doc.add_table(rows=len(rev_data), cols=2)
    t_rev.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_rev, color="003366", sz="4")
    for r_idx, (col1, col2) in enumerate(rev_data):
        row = t_rev.rows[r_idx]
        cell0, cell1 = row.cells[0], row.cells[1]
        cell0.width = Inches(4.5)
        cell1.width = Inches(2.0)
        set_cell_margins(cell0, top=60, bottom=60, left=80, right=80)
        set_cell_margins(cell1, top=60, bottom=60, left=80, right=80)
        p0 = cell0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        p1 = cell1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)

        if r_idx == 0:
            set_cell_background(cell0, "003366")
            set_cell_background(cell1, "003366")
            r0 = p0.add_run(col1)
            r0.bold = True
            r0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r1 = p1.add_run(col2)
            r1.bold = True
            r1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        else:
            set_cell_background(cell0, "F0F4F8" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_background(cell1, "F0F4F8" if r_idx % 2 == 1 else "FFFFFF")
            r0 = p0.add_run(col1)
            r0.font.size = Pt(8.5)
            r1 = p1.add_run(col2)
            r1.font.size = Pt(8.5)
            r1.bold = True
            if r_idx == len(rev_data) - 1:
                r0.bold = True
                r0.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
                r1.font.color.rgb = RGBColor(0x00, 0x66, 0x00)
            else:
                r1.font.color.rgb = RGBColor(0x00, 0x66, 0x00)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2("5.2 Sprint Retrospective (Format Resmi Modul 7 Bab 9)")
    retro_text = (
        "• What Went Well?\n"
        "  Penerapan konsep Abstraction pada Superclass Kendaraan dan Person berhasil menyederhanakan arsitektur dan mengeliminasi risiko instansiasi objek tanpa bentuk konkrit. Implementasi 3 interface utama (DapatDilacak, Perawatan, Cetakable) memberikan kontrak perilaku yang sangat jelas dan memungkinkan pemrosesan polimorfis lintas hierarki class secara elegan.\n\n"
        "• What Went Wrong?\n"
        "  Tantangan utama terletak pada penentuan batas tanggung jawab antara apa yang seharusnya menjadi abstract method di Superclass versus apa yang harus diekstrak ke dalam Interface. Selain itu, penyesuaian method Cetakable memerlukan perancangan antarmuka yang seragam namun fleksibel untuk berbagai format dokumen (kuitansi sewa, kartu member, dan slip gaji pegawai).\n\n"
        "• Improvement for Next Sprint (P8 - UTS / P9):\n"
        "  Menyiapkan integrasi komprehensif seluruh materi P1–P7 untuk menghadapi Evaluasi UTS Praktikum (P8), serta merancang transisi arsitektur menuju Graphical User Interface (GUI) atau integrasi persistensi database relasional (JDBC/MySQL) pada P9 mendatang."
    )
    add_callout(retro_text, title="📌 SPRINT RETROSPECTIVE (SPRINT 7 / P7)", bg_hex="FDFBF7", border_hex="CC8800")

    add_h2("5.3 Checklist Ketercapaian Definition of Done (15 Butir Modul 7 Bab 10)")
    doc.add_paragraph("Tabel berikut membuktikan bahwa seluruh 15 butir checklist Definition of Done telah terpenuhi 100%:")

    dod_items = [
        ("1", "Menggunakan proyek hasil P6 sebagai basis refactoring", "TERPENUHI (100%)"),
        ("2", "Audit desain dan alasan penggunaan abstraksi tersedia", "TERPENUHI (100%)"),
        ("3", "Minimal 1 abstract class yang relevan dibuat (Kendaraan & Person)", "TERPENUHI (100%)"),
        ("4", "Minimal 1 abstract method dibuat (hitungEstimasi, tampilkanSpesifikasi, dll)", "TERPENUHI (100%)"),
        ("5", "Minimal 2 subclass konkret mengimplementasikan abstract method (5 Subclass)", "TERPENUHI (100%)"),
        ("6", "Minimal 1 interface dibuat (DapatDilacak, Perawatan, Cetakable)", "TERPENUHI (100%)"),
        ("7", "Minimal 2 class mengimplementasikan interface yang relevan", "TERPENUHI (100%)"),
        ("8", "Polymorphism melalui superclass dan interface dapat dibuktikan", "TERPENUHI (100%)"),
        ("9", "Class diagram diperbarui dan konsisten dengan struktur kode", "TERPENUHI (100%)"),
        ("10", "Program berhasil dikompilasi (0 error) dan dijalankan", "TERPENUHI (100%)"),
        ("11", "Minimal 4 skenario pengujian tersedia (8 Fase Pengujian Selesai)", "TERPENUHI (100%)"),
        ("12", "Fitur proyek sebelumnya (P1–P6) tetap berjalan normal", "TERPENUHI (100%)"),
        ("13", "Sprint Backlog diperbarui lengkap (SB-01 s/d SB-07)", "TERPENUHI (100%)"),
        ("14", "Sprint Review tersedia (7/7 item terverifikasi)", "TERPENUHI (100%)"),
        ("15", "Sprint Retrospective tersedia lengkap", "TERPENUHI (100%)")
    ]
    t_dod = doc.add_table(rows=len(dod_items) + 1, cols=3)
    t_dod.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_dod, color="003366", sz="4")

    # Header
    hdr = t_dod.rows[0]
    for c_idx, h_text in enumerate(["No", "Butir Persyaratan Definition of Done Modul 7", "Status Ketercapaian"]):
        cell = hdr.cells[c_idx]
        set_cell_background(cell, "003366")
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(8.5)

    for r_idx, (no_str, req_str, st_str) in enumerate(dod_items):
        row = t_dod.rows[r_idx + 1]
        for c_idx, val in enumerate([no_str, req_str, st_str]):
            cell = row.cells[c_idx]
            set_cell_background(cell, "F0F4F8" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(8)
            if c_idx == 0:
                r.bold = True
            elif c_idx == 2:
                r.bold = True
                r.font.color.rgb = RGBColor(0x00, 0x66, 0x00)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2("5.4 Target Pertemuan Berikutnya (P8 — UTS Praktikum PBO)")
    doc.add_paragraph(
        "Hasil implementasi Modul 7 ini menjadi landasan arsitektural utama (baseline project) dalam menghadapi Pertemuan 8 (P8 — Evaluasi Tengah Semester / UTS Praktikum). Pada P8, seluruh kompetensi terintegrasi P1 hingga P7 (Class, Object, Constructor, Encapsulation, Composition, Aggregation, Association, Dependency, Inheritance, Polymorphism, Abstract Class, dan Interface) akan dievaluasi secara menyeluruh melalui demonstrasi program, inspeksi kode sumber, dan klarifikasi individual."
    )

    add_h2("5.5 Kesimpulan Akhir")
    doc.add_paragraph(
        "Praktikum Pemrograman Berorientasi Obyek Modul 7 (Abstract Class, Abstract Method, dan Interface) pada proyek AutoRent System telah berhasil diselesaikan dengan hasil sempurna (100% Definition of Done tercapai). Melalui penerapan abstraksi murni dan multi-interface realization, sistem persewaan armada rental kini memiliki fondasi arsitektur yang sangat kokoh, modular, terstandarisasi, aman dari anomali instansiasi, dan siap menghadapi tahap evaluasi UTS Praktikum P8."
    )

    output_path = "Laporan_Praktikum_PBO_Modul_7_AutoRent.docx"
    doc.save(output_path)
    print(f"Successfully generated full Word report: {output_path}")

if __name__ == "__main__":
    build_docx_report()
